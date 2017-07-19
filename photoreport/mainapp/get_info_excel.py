from os.path import join, dirname, abspath
import xlrd
import json
from docx import Document
import datetime



def read_from_excel(fname):
    dict_list = []
    #fname = join(dirname(dirname(abspath(__file__))), 'template_1', 'project-List.xls')

    # Open the workbook
    xl_workbook = xlrd.open_workbook(fname)
    print ('Imported Project List ' )
    #print fname

    # List sheet names, and pull a sheet by name
    sheet_names = xl_workbook.sheet_names()
    #print('Sheet Names', sheet_names)

    # match the "ProjectListReport" in sheet names.
    xl_sheet = ''
    for sheet_name in sheet_names:
        if sheet_name == 'ProjectListReport':
            xl_sheet = xl_workbook.sheet_by_name(sheet_name)
            #print ('Importing Sheet name: %s' % xl_sheet.name)
            break
        continue

    # read header values into the list
    keys_1 = []
    prev_key = ''
    for col_index in xrange(xl_sheet.ncols):
        if col_index >= (xl_sheet.ncols - 5):
            keys_1.append('')
        else:
            if prev_key == '' and xl_sheet.cell(0, col_index).ctype == 0:
                keys_1.append('')
                continue
            elif xl_sheet.cell(0, col_index).ctype == 1:
                if prev_key == xl_sheet.cell(0, col_index).value:
                    keys_1.append(prev_key)
                else:
                    prev_key = xl_sheet.cell(0, col_index).value
                    keys_1.append(prev_key)
            elif prev_key != '' and xl_sheet.cell(0, col_index).ctype == 0:
                keys_1.append(prev_key)

    keys_2 = [xl_sheet.cell(1, col_index).value.replace('\n', '') for col_index in xrange(xl_sheet.ncols)]

    keys = []   # table head title.
    for index in xrange(xl_sheet.ncols):
        if keys_1[index] != '':
            keys.append(keys_1[index] + '/' + keys_2[index])
            continue
        keys.append(keys_2[index])

    # print keys
    for row_index in xrange(3, xl_sheet.nrows):
        #print 'row_index: {}'.format(row_index)
        d = {}
        for col_index in xrange(xl_sheet.ncols):
            if xl_sheet.cell(row_index, col_index).ctype == 3:
                time_excel = str(datetime.datetime(*xlrd.xldate_as_tuple(xl_sheet.cell(row_index, col_index).value,
                                                                              xl_workbook.datemode)))[:10]
                d.update({keys[col_index]: time_excel})
            else:
                if xl_sheet.cell(row_index, col_index).value == 'TX':
                    d.update({keys[col_index]: 'Texas'})
                else:
                    d.update({keys[col_index]: xl_sheet.cell(row_index, col_index).value})
        dict_list.append(d)

    #print dict_list #debugging imported data print
    return dict_list










